"""Creates sample data files for testing the ERP Data Entry workflow."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def create_sample_erp_data():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # Header row
    headers = ["invoice_number", "vendor_name", "invoice_date", "amount",
               "description", "cost_center", "purchase_order"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Sample data rows
    data = [
        ["INV-2026-001", "Acme Corporation",    "2026-01-15", 1500.00, "Office supplies",     "CC-001", "PO-5001"],
        ["INV-2026-002", "Beta Technologies",   "2026-01-22", 8750.50, "Software licenses",   "CC-002", "PO-5002"],
        ["INV-2026-003", "Gamma Logistics",     "2026-02-03", 3200.00, "Freight services",    "CC-003", "PO-5003"],
        ["INV-2026-004", "Delta Consulting",    "2026-02-14", 12000.00,"Consulting fees",     "CC-001", "PO-5004"],
        ["INV-2026-005", "Epsilon Supplies",    "2026-02-28", 450.75,  "Stationery",          "CC-004", "PO-5005"],
        ["INV-2026-006", "Zeta Engineering",    "2026-03-05", 6800.00, "Equipment repair",    "CC-002", "PO-5006"],
        ["INV-2026-007", "Eta Media Group",     "2026-03-12", 2100.00, "Marketing materials", "CC-005", "PO-5007"],
        ["INV-2026-008", "Theta IT Solutions",  "2026-03-20", 9500.00, "IT infrastructure",   "CC-002", "PO-5008"],
        ["INV-2026-009", "Iota Catering",       "2026-04-01", 875.00,  "Event catering",      "CC-006", "PO-5009"],
        ["INV-2026-010", "Kappa Legal Services","2026-04-15", 4200.00, "Legal advisory",      "CC-001", "PO-5010"],
    ]

    for row_data in data:
        ws.append(row_data)

    # Auto-size columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    os.makedirs("data", exist_ok=True)
    path = "data/sample_invoices.xlsx"
    wb.save(path)
    print(f"Sample data created: {path} ({len(data)} rows)")
    return path

if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    create_sample_erp_data()
